"""The company 'Tool Verification' workbook: read tool rows, write updates.

Reading uses openpyxl (read-only, values only). Writing goes through
`XlsxPackage` so the rest of the file stays exactly as it was.

Expected layout per regional sheet (header names, not fixed letters):

    Cal Next | Cal Last | Manufacturer | Model | Description | Serial |
    Asset # | Cal Interval | Cal Cycle | Label Photo | SMAX Tool Photo
"""

from __future__ import annotations

import datetime as dt
import io
import os
import posixpath
import re
import shutil
from dataclasses import dataclass, field

import openpyxl
from lxml import etree
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils import column_index_from_string, get_column_letter

from .xlsx_package import (
    CT_DRAWING,
    NS,
    REL_DRAWING,
    REL_IMAGE,
    REL_STYLES,
    REL_WORKSHEET,
    XlsxPackage,
    q,
)

# Header aliases (lower-cased, spaces collapsed) -> logical column key.
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "cal_next": ("cal next", "next cal", "cal due", "due date", "차기교정일", "차기 교정일"),
    "cal_last": ("cal last", "last cal", "cal date", "교정일"),
    "manufacturer": ("manufacturer", "제조사"),
    "model": ("model", "모델"),
    "description": ("description", "설명", "품명"),
    "serial": ("serial", "serial #", "serial no", "s/n", "시리얼", "시리얼번호"),
    "asset": ("asset #", "asset", "asset no", "자산번호", "담당자"),
    "interval": ("cal interval", "interval", "교정주기"),
    "cycle": ("cal cycle", "cycle", "주기단위"),
    "label_photo": ("label photo", "라벨 사진", "교정필증"),
    "smax_photo": ("smax tool photo", "smax photo", "smax"),
}
REQUIRED_COLUMNS = ("cal_next", "cal_last", "serial", "asset")

# East-Asian locale built-in date formats that openpyxl does not list.
CJK_DATE_FORMAT_IDS = set(range(27, 37)) | set(range(50, 59))

MIN_PICTURE_ROW_PT = 60.0  # rows lower than this are raised before a picture goes in
RAISED_ROW_PT = 100.0
PICTURE_PAD_PX = 3
EMU_PER_PX = 9525
EMU_PER_PT = 12700
MAX_DIGIT_WIDTH_PX = 7  # Calibri 11, the default font of the company file

HANGUL_NAME = re.compile(r"[가-힣]{2,4}")
NON_NAME_WORDS = {
    "공통", "등록", "메디아나", "사무실", "폐기", "예비", "서울", "부산", "대구", "인천", "광주",
    "대전", "울산", "세종", "경기", "강원", "충북", "충남", "전북", "전남", "경북", "경남", "제주",
}


class WorkbookError(Exception):
    pass


def normalize_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def serial_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_cell_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        m = re.search(r"(\d{4})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})", value)
        if m:
            try:
                return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                return None
    return None


def interval_in_months(interval, cycle) -> int | None:
    try:
        n = int(float(str(interval).strip()))
    except (TypeError, ValueError):
        return None
    unit = str(cycle or "MONTHS").strip().upper()
    if unit.startswith("Y") or unit.startswith("년"):
        return n * 12
    if unit.startswith("M") or unit.startswith("개월") or unit.startswith("월"):
        return n
    return None


def extract_names(asset: str, exclude: set[str] = frozenset()) -> list[str]:
    """Engineer names in an 'Asset #' value such as 'South_광주_CT_강건욱님 - 사무실'."""
    names = []
    for token in HANGUL_NAME.findall(asset):
        if token.endswith("님"):
            token = token[:-1]
        if len(token) >= 2 and token not in NON_NAME_WORDS and token not in exclude and token not in names:
            names.append(token)
    return names


@dataclass
class ToolRecord:
    sheet: str
    row: int
    cal_next: dt.date | None
    cal_next_raw: object
    cal_last: dt.date | None
    manufacturer: str
    model: str
    description: str
    serial: str
    asset: str
    interval_months: int | None
    interval_raw: str = ""
    names: list[str] = field(default_factory=list)
    picture_counts: dict[str, int] = field(default_factory=dict)

    @property
    def key(self) -> str:
        return f"{self.sheet}!{self.row}"

    @property
    def disposed(self) -> bool:
        return isinstance(self.cal_next_raw, str) and "폐기" in self.cal_next_raw

    def label(self) -> str:
        return f"[{self.sheet} {self.row}행] {self.asset} · {self.model} · S/N {self.serial}"


@dataclass
class SheetInfo:
    name: str
    part: str
    header_row: int
    columns: dict[str, str]  # logical key -> column letter


class ToolWorkbook:
    def __init__(self, data: bytes, path: str | None = None):
        self.path = path
        self.pkg = XlsxPackage(data)
        self.sheets: dict[str, SheetInfo] = {}
        self.records: list[ToolRecord] = []
        self._styles_part = None
        self._date1904 = False
        self._map_parts()
        self._read_values(data)
        self._count_pictures()

    @classmethod
    def open(cls, path: str | os.PathLike) -> "ToolWorkbook":
        with open(path, "rb") as fh:
            return cls(fh.read(), str(path))

    # ------------------------------------------------------------------ read
    def _map_parts(self) -> None:
        wb_part = "xl/workbook.xml"
        root = self.pkg.xml(wb_part)
        pr = root.find(q("main", "workbookPr"))
        self._date1904 = pr is not None and pr.get("date1904") in ("1", "true")
        targets = self.pkg.rel_targets(wb_part)
        self._sheet_parts: dict[str, str] = {}
        for sheet in root.iter(q("main", "sheet")):
            rid = sheet.get(q("r", "id"))
            rtype, part = targets.get(rid, (None, None))
            if rtype == REL_WORKSHEET:
                self._sheet_parts[sheet.get("name")] = part
        for rtype, part in targets.values():
            if rtype == REL_STYLES:
                self._styles_part = part

    def _read_values(self, data: bytes) -> None:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                if ws.title not in self._sheet_parts:
                    continue
                rows = list(ws.iter_rows(values_only=True))
                header_idx, columns = self._find_header(rows)
                if header_idx is None:
                    continue
                info = SheetInfo(ws.title, self._sheet_parts[ws.title], header_idx + 1, columns)
                self.sheets[ws.title] = info
                for offset, values in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
                    rec = self._make_record(info, offset, values)
                    if rec is not None:
                        self.records.append(rec)
        finally:
            wb.close()
        for rec in self.records:
            rec.names = extract_names(rec.asset, exclude=set(self.sheets))
        if not self.sheets:
            raise WorkbookError(
                "툴 목록 시트를 찾지 못했습니다. 헤더에 'Serial', 'Asset #', 'Cal Next', 'Cal Last' 열이 있어야 합니다."
            )

    @staticmethod
    def _find_header(rows) -> tuple[int | None, dict[str, str]]:
        for idx, values in enumerate(rows[:10]):
            columns: dict[str, str] = {}
            for col, value in enumerate(values or (), start=1):
                header = normalize_header(value)
                for key, aliases in HEADER_ALIASES.items():
                    if key not in columns and header in aliases:
                        columns[key] = get_column_letter(col)
            if all(k in columns for k in REQUIRED_COLUMNS):
                return idx, columns
        return None, {}

    @staticmethod
    def _make_record(info: SheetInfo, row: int, values) -> ToolRecord | None:
        def get(key):
            letter = info.columns.get(key)
            if letter is None:
                return None
            i = column_index_from_string(letter) - 1
            return values[i] if i < len(values) else None

        serial = serial_to_text(get("serial"))
        asset = str(get("asset") or "").strip()
        if not serial and not asset:
            return None
        cal_next_raw = get("cal_next")
        return ToolRecord(
            sheet=info.name,
            row=row,
            cal_next=parse_cell_date(cal_next_raw),
            cal_next_raw=cal_next_raw,
            cal_last=parse_cell_date(get("cal_last")),
            manufacturer=str(get("manufacturer") or "").strip(),
            model=serial_to_text(get("model")),
            description=str(get("description") or "").strip(),
            serial=serial,
            asset=asset,
            interval_months=interval_in_months(get("interval"), get("cycle")),
            interval_raw=f"{get('interval') or ''} {get('cycle') or ''}".strip(),
        )

    def _count_pictures(self) -> None:
        for rec in self.records:
            info = self.sheets[rec.sheet]
            for key in ("label_photo", "smax_photo"):
                letter = info.columns.get(key)
                if letter:
                    rec.picture_counts[key] = len(self._anchors_at(info, rec.row, letter))

    def find_record(self, sheet: str, row: int) -> ToolRecord | None:
        for rec in self.records:
            if rec.sheet == sheet and rec.row == row:
                return rec
        return None

    # ------------------------------------------------------------ cell writes
    def _sheet_root(self, info: SheetInfo) -> etree._Element:
        return self.pkg.xml(info.part)

    def _row_el(self, info: SheetInfo, row: int, create: bool = True) -> etree._Element | None:
        sheet_data = self._sheet_root(info).find(q("main", "sheetData"))
        for el in sheet_data.findall(q("main", "row")):
            r = int(el.get("r"))
            if r == row:
                return el
            if r > row:
                if not create:
                    return None
                new = etree.Element(q("main", "row"))
                new.set("r", str(row))
                el.addprevious(new)
                return new
        if not create:
            return None
        new = etree.SubElement(sheet_data, q("main", "row"))
        new.set("r", str(row))
        return new

    def _cell_el(self, info: SheetInfo, letter: str, row: int) -> etree._Element:
        row_el = self._row_el(info, row)
        ref = f"{letter}{row}"
        target_idx = column_index_from_string(letter)
        for c in row_el.findall(q("main", "c")):
            m = re.match(r"([A-Z]+)", c.get("r", ""))
            idx = column_index_from_string(m.group(1)) if m else 0
            if idx == target_idx:
                return c
            if idx > target_idx:
                new = etree.Element(q("main", "c"))
                new.set("r", ref)
                c.addprevious(new)
                return new
        new = etree.SubElement(row_el, q("main", "c"))
        new.set("r", ref)
        return new

    def _date_serial(self, value: dt.date) -> int:
        base = dt.date(1904, 1, 1) if self._date1904 else dt.date(1899, 12, 30)
        return (value - base).days

    def _style_is_date(self, s: str | None) -> bool:
        if s is None or self._styles_part is None:
            return False
        styles = self.pkg.xml(self._styles_part)
        xfs = styles.find(q("main", "cellXfs"))
        if xfs is None:
            return False
        xf_list = xfs.findall(q("main", "xf"))
        try:
            xf = xf_list[int(s)]
        except (ValueError, IndexError):
            return False
        fmt_id = int(xf.get("numFmtId", "0"))
        if fmt_id in CJK_DATE_FORMAT_IDS:
            return True
        if fmt_id in BUILTIN_FORMATS:
            return is_date_format(BUILTIN_FORMATS[fmt_id])
        num_fmts = styles.find(q("main", "numFmts"))
        if num_fmts is not None:
            for nf in num_fmts.findall(q("main", "numFmt")):
                if int(nf.get("numFmtId", "-1")) == fmt_id:
                    return is_date_format(nf.get("formatCode", ""))
        return False

    def _date_style_for(self, info: SheetInfo, letter: str) -> str:
        """A date cell style: from the same column if possible, else any, else a new one."""
        sheet_data = self._sheet_root(info).find(q("main", "sheetData"))
        for c in sheet_data.iter(q("main", "c")):
            if re.match(rf"^{letter}\d+$", c.get("r", "")) and self._style_is_date(c.get("s")):
                return c.get("s")
        styles = self.pkg.xml(self._styles_part)
        xfs = styles.find(q("main", "cellXfs"))
        for i in range(len(xfs.findall(q("main", "xf")))):
            if self._style_is_date(str(i)):
                return str(i)
        xf = etree.SubElement(xfs, q("main", "xf"))
        for k, v in (("numFmtId", "14"), ("fontId", "0"), ("fillId", "0"), ("borderId", "0"),
                     ("xfId", "0"), ("applyNumberFormat", "1")):
            xf.set(k, v)
        count = len(xfs.findall(q("main", "xf")))
        xfs.set("count", str(count))
        self.pkg.touch(self._styles_part)
        return str(count - 1)

    def write_date(self, rec: ToolRecord, key: str, value: dt.date) -> str:
        """Write a date into a tool row. Returns 'written', 'unchanged' or 'formula'."""
        info = self.sheets[rec.sheet]
        letter = info.columns[key]
        c = self._cell_el(info, letter, rec.row)
        if c.find(q("main", "f")) is not None:
            # Leave formulas alone (e.g. Cal Next = EDATE(Cal Last, 12)); make
            # Excel recompute them when the file is opened.
            self._force_full_calc()
            return "formula"
        serial = str(self._date_serial(value))
        v = c.find(q("main", "v"))
        if c.get("t") is None and v is not None and v.text == serial and self._style_is_date(c.get("s")):
            return "unchanged"
        for child in list(c):
            c.remove(child)
        if "t" in c.attrib:
            del c.attrib["t"]
        if not self._style_is_date(c.get("s")):
            c.set("s", self._date_style_for(info, letter))
        etree.SubElement(c, q("main", "v")).text = serial
        self.pkg.touch(info.part)
        if key == "cal_next":
            rec.cal_next, rec.cal_next_raw = value, dt.datetime.combine(value, dt.time())
        elif key == "cal_last":
            rec.cal_last = value
        return "written"

    def _force_full_calc(self) -> None:
        root = self.pkg.xml("xl/workbook.xml")
        calc = root.find(q("main", "calcPr"))
        if calc is None:
            calc = etree.SubElement(root, q("main", "calcPr"))
        calc.set("fullCalcOnLoad", "1")
        self.pkg.touch("xl/workbook.xml")

    # -------------------------------------------------------------- geometry
    def _col_width_px(self, info: SheetInfo, col_idx: int) -> int:
        root = self._sheet_root(info)
        cols = root.find(q("main", "cols"))
        if cols is not None:
            for col in cols.findall(q("main", "col")):
                if int(col.get("min")) <= col_idx <= int(col.get("max")):
                    if col.get("hidden") in ("1", "true"):
                        return 0
                    width = float(col.get("width", "8.43"))
                    return int(((256 * width + int(128 / MAX_DIGIT_WIDTH_PX)) / 256) * MAX_DIGIT_WIDTH_PX)
        fmt = root.find(q("main", "sheetFormatPr"))
        if fmt is not None and fmt.get("defaultColWidth"):
            return int(float(fmt.get("defaultColWidth")) * MAX_DIGIT_WIDTH_PX + 5)
        return 64

    def _default_row_pt(self, info: SheetInfo) -> float:
        fmt = self._sheet_root(info).find(q("main", "sheetFormatPr"))
        if fmt is not None and fmt.get("defaultRowHeight"):
            return float(fmt.get("defaultRowHeight"))
        return 15.0

    def _row_height_pt(self, info: SheetInfo, row: int) -> float:
        el = self._row_el(info, row, create=False)
        if el is not None:
            if el.get("hidden") in ("1", "true"):
                return 0.0
            if el.get("ht"):
                return float(el.get("ht"))
        return self._default_row_pt(info)

    def _ensure_picture_row_height(self, info: SheetInfo, row: int) -> float:
        height = self._row_height_pt(info, row)
        if height < MIN_PICTURE_ROW_PT:
            el = self._row_el(info, row)
            el.set("ht", str(RAISED_ROW_PT))
            el.set("customHeight", "1")
            self.pkg.touch(info.part)
            height = RAISED_ROW_PT
        return height

    def _abs_offset_emu(self, info: SheetInfo, col_idx: int, row: int) -> tuple[int, int]:
        x = sum(self._col_width_px(info, c) for c in range(1, col_idx)) * EMU_PER_PX
        heights = {}
        sheet_data = self._sheet_root(info).find(q("main", "sheetData"))
        for el in sheet_data.findall(q("main", "row")):
            heights[int(el.get("r"))] = el
        default = self._default_row_pt(info)
        y_pt = 0.0
        for r in range(1, row):
            el = heights.get(r)
            if el is not None and el.get("hidden") in ("1", "true"):
                continue
            y_pt += float(el.get("ht")) if el is not None and el.get("ht") else default
        return x, int(y_pt * EMU_PER_PT)

    # -------------------------------------------------------------- pictures
    def _drawing_part(self, info: SheetInfo, create: bool) -> str | None:
        root = self._sheet_root(info)
        drawing = root.find(q("main", "drawing"))
        if drawing is not None:
            rtype, part = self.pkg.rel_targets(info.part).get(drawing.get(q("r", "id")), (None, None))
            if rtype == REL_DRAWING:
                return part
        if not create:
            return None
        part = self.pkg.unique_part_name("xl/drawings", "drawing", "xml")
        ws_dr = etree.Element(q("xdr", "wsDr"), nsmap={"xdr": NS["xdr"], "a": NS["a"]})
        self.pkg.write_xml(part, ws_dr)
        self.pkg.add_override(part, CT_DRAWING)
        rid = self.pkg.add_rel(info.part, REL_DRAWING, part)
        el = etree.Element(q("main", "drawing"), nsmap={"r": NS["r"]})
        el.set(q("r", "id"), rid)
        follows = ("legacyDrawing", "legacyDrawingHF", "drawingHF", "picture", "oleObjects",
                   "controls", "webPublishItems", "tableParts", "extLst")
        for child in root:
            if isinstance(child.tag, str) and etree.QName(child).localname in follows:
                child.addprevious(el)
                break
        else:
            root.append(el)
        self.pkg.touch(info.part)
        return part

    def _anchors_at(self, info: SheetInfo, row: int, letter: str) -> list[etree._Element]:
        part = self._drawing_part(info, create=False)
        if part is None:
            return []
        col0 = column_index_from_string(letter) - 1
        row0 = row - 1
        found = []
        for anchor in self.pkg.xml(part):
            if not isinstance(anchor.tag, str):
                continue
            if etree.QName(anchor).localname not in ("twoCellAnchor", "oneCellAnchor"):
                continue
            if anchor.find(q("xdr", "pic")) is None:
                continue
            frm = anchor.find(q("xdr", "from"))
            if frm is None:
                continue
            if int(frm.findtext(q("xdr", "col"))) == col0 and int(frm.findtext(q("xdr", "row"))) == row0:
                found.append(anchor)
        return found

    def replace_picture(
        self,
        rec: ToolRecord,
        key: str,
        data: bytes,
        ext: str,
        size_px: tuple[int, int],
        name: str = "",
        descr: str = "",
    ) -> int:
        """Put a picture in the row's `key` column, replacing pictures already there.

        Returns how many old pictures were removed.
        """
        info = self.sheets[rec.sheet]
        letter = info.columns.get(key)
        if letter is None:
            raise WorkbookError(f"'{rec.sheet}' 시트에 사진 열({key})이 없습니다.")
        drawing = self._drawing_part(info, create=True)
        ws_dr = self.pkg.xml(drawing)

        removed = self._anchors_at(info, rec.row, letter)
        old_rids = set()
        for anchor in removed:
            for blip in anchor.iter(q("a", "blip")):
                if blip.get(q("r", "embed")):
                    old_rids.add(blip.get(q("r", "embed")))
            ws_dr.remove(anchor)
        if removed:
            self.pkg.touch(drawing)
            self._drop_unused_images(drawing, old_rids)

        ext = ext.lower().lstrip(".")
        media = self.pkg.unique_part_name("xl/media", "image", ext)
        self.pkg.write(media, data)
        self.pkg.ensure_default_content_type(ext, "image/jpeg" if ext in ("jpg", "jpeg") else f"image/{ext}")
        rid = self.pkg.add_rel(drawing, REL_IMAGE, media)

        col_idx = column_index_from_string(letter)
        row_pt = self._ensure_picture_row_height(info, rec.row)
        box_w = max(self._col_width_px(info, col_idx) - 2 * PICTURE_PAD_PX, 10)
        box_h = max(int(row_pt * 96 / 72) - 2 * PICTURE_PAD_PX, 10)
        img_w, img_h = size_px
        scale = min(box_w / img_w, box_h / img_h)
        cx = int(img_w * scale) * EMU_PER_PX
        cy = int(img_h * scale) * EMU_PER_PX
        pad = PICTURE_PAD_PX * EMU_PER_PX
        abs_x, abs_y = self._abs_offset_emu(info, col_idx, rec.row)

        ws_dr.append(
            self._anchor_xml(
                col0=col_idx - 1,
                row0=rec.row - 1,
                pad=pad,
                cx=cx,
                cy=cy,
                abs_xy=(abs_x + pad, abs_y + pad),
                shape_id=self._next_shape_id(ws_dr),
                rid=rid,
                name=name or f"{key} {rec.serial}",
                descr=descr,
            )
        )
        self.pkg.touch(drawing)
        rec.picture_counts[key] = 1
        return len(removed)

    def _next_shape_id(self, ws_dr: etree._Element) -> int:
        ids = [int(el.get("id")) for el in ws_dr.iter(q("xdr", "cNvPr")) if (el.get("id") or "").isdigit()]
        return max(ids, default=1) + 1

    def _drop_unused_images(self, drawing: str, rids: set[str]) -> None:
        still_used = set()
        for el in self.pkg.xml(drawing).iter():
            for attr in (q("r", "embed"), q("r", "link"), q("r", "id")):
                if el.get(attr):
                    still_used.add(el.get(attr))
        targets = self.pkg.rel_targets(drawing)
        for rid in rids - still_used:
            _rtype, media = targets.get(rid, (None, None))
            self.pkg.remove_rel(drawing, rid)
            if media and self.pkg.exists(media) and not self.pkg.is_referenced(media):
                self.pkg.delete(media)

    @staticmethod
    def _anchor_xml(col0, row0, pad, cx, cy, abs_xy, shape_id, rid, name, descr) -> etree._Element:
        xdr = NS["xdr"]
        a = NS["a"]
        anchor = etree.Element(f"{{{xdr}}}twoCellAnchor", nsmap={"xdr": xdr, "a": a})
        anchor.set("editAs", "oneCell")
        for tag, col_off, row_off in (("from", pad, pad), ("to", pad + cx, pad + cy)):
            pos = etree.SubElement(anchor, f"{{{xdr}}}{tag}")
            etree.SubElement(pos, f"{{{xdr}}}col").text = str(col0)
            etree.SubElement(pos, f"{{{xdr}}}colOff").text = str(col_off)
            etree.SubElement(pos, f"{{{xdr}}}row").text = str(row0)
            etree.SubElement(pos, f"{{{xdr}}}rowOff").text = str(row_off)
        pic = etree.SubElement(anchor, f"{{{xdr}}}pic")
        nv = etree.SubElement(pic, f"{{{xdr}}}nvPicPr")
        c_nv_pr = etree.SubElement(nv, f"{{{xdr}}}cNvPr")
        c_nv_pr.set("id", str(shape_id))
        c_nv_pr.set("name", name)
        if descr:
            c_nv_pr.set("descr", descr)
        c_nv_pic_pr = etree.SubElement(nv, f"{{{xdr}}}cNvPicPr")
        etree.SubElement(c_nv_pic_pr, f"{{{a}}}picLocks").set("noChangeAspect", "1")
        fill = etree.SubElement(pic, f"{{{xdr}}}blipFill")
        blip = etree.SubElement(fill, f"{{{a}}}blip", nsmap={"r": NS["r"]})
        blip.set(f"{{{NS['r']}}}embed", rid)
        etree.SubElement(etree.SubElement(fill, f"{{{a}}}stretch"), f"{{{a}}}fillRect")
        sp_pr = etree.SubElement(pic, f"{{{xdr}}}spPr")
        xfrm = etree.SubElement(sp_pr, f"{{{a}}}xfrm")
        off = etree.SubElement(xfrm, f"{{{a}}}off")
        off.set("x", str(abs_xy[0]))
        off.set("y", str(abs_xy[1]))
        ext = etree.SubElement(xfrm, f"{{{a}}}ext")
        ext.set("cx", str(cx))
        ext.set("cy", str(cy))
        geom = etree.SubElement(sp_pr, f"{{{a}}}prstGeom")
        geom.set("prst", "rect")
        etree.SubElement(geom, f"{{{a}}}avLst")
        etree.SubElement(anchor, f"{{{xdr}}}clientData")
        return anchor

    # ------------------------------------------------------------------ save
    def to_bytes(self) -> bytes:
        return self.pkg.to_bytes()

    def save(self, path: str | os.PathLike | None = None, backup_dir: str | None = None) -> str | None:
        """Save (atomically). Returns the backup path if a backup was made."""
        target = str(path or self.path)
        if not target:
            raise WorkbookError("저장할 경로가 없습니다.")
        backup = None
        if backup_dir and os.path.exists(target):
            os.makedirs(backup_dir, exist_ok=True)
            stem, ext = os.path.splitext(os.path.basename(target))
            stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
            backup = os.path.join(backup_dir, f"{stem}_{stamp}{ext}")
            shutil.copy2(target, backup)
        try:
            self.pkg.save(target)
        except PermissionError as exc:
            raise WorkbookError(
                f"엑셀 파일에 쓸 수 없습니다: {target}\n파일이 Excel에서 열려 있다면 닫고 다시 시도하세요."
            ) from exc
        self.path = target
        return backup


def media_parts(pkg: XlsxPackage) -> list[str]:
    return [n for n in pkg.names() if posixpath.dirname(n) == "xl/media"]
