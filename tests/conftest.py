import datetime as dt
import io
import os
import sys

import pytest
from PIL import Image, ImageDraw

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

HEADERS = ["Cal Next", "Cal Last", "Manufacturer", "Model", "Description", "Serial", "Asset #",
           "Cal Interval", "Cal Cycle", "Label Photo", "SMAX Tool Photo"]

# Synthetic data shaped like the company file (names/serials are made up).
SHEETS = {
    "대전": [
        (dt.datetime(2026, 12, 29), dt.datetime(2025, 12, 29), "FLUKE", "117", "TRUE RMS MULTIMETER",
         "44350259SV", "South_대전_MR_홍길동", "12", "MONTHS"),
        (dt.datetime(2026, 12, 25), dt.datetime(2026, 6, 25), "SNAP-ON", "CTECH2R100A", "TORQUE WRENCH",
         218800584, "South_대전_CT_김철수", "6", "MONTHS"),
        (dt.datetime(2027, 1, 5), dt.datetime(2026, 1, 5), "FLUKE", "287", "TRUE RMS MULTIMETER",
         46900004, "South_대전_CT_김철수", "12", "MONTHS"),
    ],
    "광주": [
        (dt.datetime(2026, 12, 29), dt.datetime(2025, 12, 29), "FLUKE", "287", "TRUE RMS MULTIMETER",
         "31550069", "South_광주_CT_이영희", "12", "MONTHS"),
        (dt.datetime(2026, 12, 29), dt.datetime(2025, 12, 29), "FLUKE", "287", "TRUE RMS MULTIMETER",
         "31550075", "South_광주_CT_박민수", "12", "MONTHS"),
        (dt.datetime(2026, 6, 24), dt.datetime(2025, 12, 24), "SNAP-ON", "CTECH2R100A", "TORQUE WRENCH",
         "0716805829", "South_광주_CT_박민수", "6", "MONTHS"),
        ("폐기", dt.datetime(2025, 6, 27), "CDI", "752MFRMH", "TORQUE WRENCH", "0116503968",
         "South_광주_CT_박민수님 - 사무실", "6", "MONTHS"),
    ],
}


def _png(color, size=(200, 120)) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    buf.seek(0)
    return buf


def build_workbook(path, with_pictures=True, third_sheet_without_drawing=False):
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        for row in rows:
            ws.append(list(row))
        for r in range(2, len(rows) + 2):
            ws.row_dimensions[r].height = 100
            for col in "AB":
                if isinstance(ws[f"{col}{r}"].value, dt.datetime):
                    ws[f"{col}{r}"].number_format = "yyyy-mm-dd"
            if with_pictures:
                ws.add_image(XLImage(_png((200, 30, 30))), f"J{r}")
                ws.add_image(XLImage(_png((30, 30, 200), (400, 120))), f"K{r}")
        ws.column_dimensions["J"].width = 32.66
        ws.column_dimensions["K"].width = 71.4
    if third_sheet_without_drawing:
        ws = wb.create_sheet("부산")
        ws.append(HEADERS)
        ws.append([dt.datetime(2026, 11, 1), dt.datetime(2025, 11, 1), "FLUKE", "115", "MULTIMETER",
                   "77777777", "South_부산_XR_최민호", "12", "MONTHS"])
        ws["A2"].number_format = "yyyy-mm-dd"
        ws["B2"].number_format = "yyyy-mm-dd"
    wb.save(path)
    return path


@pytest.fixture
def workbook_path(tmp_path):
    return str(build_workbook(tmp_path / "tools.xlsx"))


def make_tool_photo(size=(900, 1200), label_box=(250, 700, 650, 966), angle=0.0) -> Image.Image:
    """Yellow tool body, light-blue sticky note, white calibration label with dark text."""
    img = Image.new("RGB", size, (214, 160, 40))
    draw = ImageDraw.Draw(img)
    draw.rectangle((150, 150, 750, 620), fill=(160, 185, 215))  # sticky note
    draw.line((250, 300, 600, 300), fill=(40, 40, 60), width=8)
    l, t, r, b = label_box
    draw.rectangle((l, t, r, b), fill=(236, 236, 234))
    for y in range(t + 30, b - 20, 45):
        draw.line((l + 20, y, r - 20, y), fill=(30, 30, 30), width=3)
        draw.rectangle((l + 150, y + 8, r - 40, y + 28), fill=(40, 40, 40))
    draw.ellipse((380, 1020, 520, 1160), fill=(20, 20, 20))
    if angle:
        img = img.rotate(angle, expand=True, fillcolor=(60, 60, 60))
    return img


def jpeg_bytes(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=90)
    return buf.getvalue()
