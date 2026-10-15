import datetime as dt
import os

import openpyxl
from PIL import Image

from conftest import jpeg_bytes, make_tool_photo
from toolcal.pipeline import analyze_photos, applicable, apply_jobs, due_status, evaluate_all, set_target
from toolcal.recognize import Recognition, RecognitionError
from toolcal.workbook import ToolWorkbook

TODAY = dt.date.today()
NEW_LAST = TODAY - dt.timedelta(days=2)


def _add_months(d, n):
    from toolcal.matcher import add_months
    return add_months(d, n)


class StubRecognizer:
    """Returns canned readings keyed by the photo's width."""

    def __init__(self, by_width):
        self.by_width = by_width

    def recognize(self, img, known_names=None):
        value = self.by_width[img.width]
        if isinstance(value, Exception):
            raise value
        return value


def _photo(width):
    return jpeg_bytes(make_tool_photo(size=(width, 1200)))


def _tool(name, serial, last=NEW_LAST, months=12, nxt=None):
    return Recognition(image_type="tool_photo", engineer_name=name, serial_number=serial,
                       calibration_date=last.isoformat() if last else "",
                       next_due_date=(nxt or _add_months(last, months)).isoformat() if (last or nxt) else "",
                       confidence="high")


def test_end_to_end(workbook_path, tmp_path):
    wb = ToolWorkbook.open(workbook_path)
    smax = Image.new("RGB", (1500, 450), (255, 255, 255))
    recognizer = StubRecognizer({
        900: _tool("박민수", "31550075"),                     # clean match
        901: _tool("김철수", "218800584", months=6),           # torque wrench, 6 months
        902: _tool("홍길동", "99999999"),                      # unknown serial
        903: RecognitionError("Claude API에 연결할 수 없습니다."),
        1500: Recognition(image_type="smax_screenshot", serial_number="31550075",
                          next_due_date=_add_months(NEW_LAST, 12).isoformat()),
    })
    files = [("a.jpg", _photo(900)), ("b.jpg", _photo(901)), ("c.jpg", _photo(902)),
             ("d.jpg", _photo(903)), ("smax.png", jpeg_bytes(smax))]
    jobs = analyze_photos(wb, files, recognizer)
    by_name = {j.filename: j for j in jobs}
    assert [j.filename for j in jobs] == ["a.jpg", "b.jpg", "c.jpg", "d.jpg", "smax.png"]
    assert by_name["a.jpg"].status == "ok" and by_name["a.jpg"].include
    assert by_name["a.jpg"].label_crop is not None
    assert by_name["b.jpg"].status == "ok"
    assert by_name["c.jpg"].status == "error" and not applicable(by_name["c.jpg"])
    assert by_name["d.jpg"].status == "error"
    assert by_name["smax.png"].kind == "smax_screenshot" and by_name["smax.png"].status == "ok"

    archive = str(tmp_path / "archive")
    changes = apply_jobs(wb, jobs, archive)
    assert len(changes) == 3
    out = tmp_path / "out.xlsx"
    wb.save(out, backup_dir=str(tmp_path / "backup"))

    ws = openpyxl.load_workbook(out)["광주"]
    assert ws["B3"].value.date() == NEW_LAST
    assert ws["A3"].value.date() == _add_months(NEW_LAST, 12)
    ws = openpyxl.load_workbook(out)["대전"]
    assert ws["B3"].value.date() == NEW_LAST
    assert ws["A3"].value.date() == _add_months(NEW_LAST, 6)

    again = ToolWorkbook.open(out)
    rec = next(r for r in again.records if r.serial == "31550075")
    assert rec.picture_counts == {"label_photo": 1, "smax_photo": 1}
    archived = os.listdir(os.path.join(archive, "광주"))
    assert any(n.startswith(f"{NEW_LAST.isoformat()}_박민수_31550075") for n in archived)
    assert any(n.endswith("_SMAX.jpg") for n in archived)
    assert os.path.exists(os.path.join(archive, "update_log.csv"))


def test_manual_target_after_failed_recognition(workbook_path):
    wb = ToolWorkbook.open(workbook_path)
    jobs = analyze_photos(wb, [("x.jpg", _photo(900))], StubRecognizer({900: RecognitionError("down")}))
    job = jobs[0]
    assert job.status == "error"
    target = next(r for r in wb.records if r.serial == "46900004")
    set_target(job, target)
    assert job.status == "review"  # dates still missing
    job.cal_last, job.cal_next = NEW_LAST, None
    evaluate_all(jobs)
    assert job.cal_next == _add_months(NEW_LAST, 12)  # derived from the 12-month interval
    job.include = True
    assert applicable(job)


def test_checks_that_need_review(workbook_path):
    wb = ToolWorkbook.open(workbook_path)
    recognizer = StubRecognizer({
        900: _tool("박민수", "31550075", nxt=_add_months(NEW_LAST, 6)),       # interval mismatch
        901: _tool("이영희", "31550069", last=dt.date(2024, 1, 2)),           # older than Excel
        902: _tool("홍길동", "44350259SV", last=TODAY + dt.timedelta(days=30)),  # future date
    })
    jobs = analyze_photos(wb, [(f"{w}.jpg", _photo(w)) for w in (900, 901, 902)], recognizer)
    assert [j.status for j in jobs] == ["review", "review", "review"]
    assert not any(j.include for j in jobs)
    assert any("교정주기" in m for m in jobs[0].messages)
    assert any("이전 날짜" in m for m in jobs[1].messages)
    assert any("미래" in m for m in jobs[2].messages)


def test_duplicate_photos_need_review(workbook_path):
    wb = ToolWorkbook.open(workbook_path)
    reading = _tool("박민수", "31550075")
    jobs = analyze_photos(wb, [("1.jpg", _photo(900)), ("2.jpg", _photo(900))], StubRecognizer({900: reading}))
    assert all(j.status == "review" for j in jobs)


def test_row_shift_is_detected(workbook_path, tmp_path):
    wb = ToolWorkbook.open(workbook_path)
    jobs = analyze_photos(wb, [("a.jpg", _photo(900))], StubRecognizer({900: _tool("박민수", "31550075")}))
    # someone inserts a row above the tool before the update is applied
    xl = openpyxl.load_workbook(workbook_path)
    xl["광주"].insert_rows(2)
    xl.save(workbook_path)
    fresh = ToolWorkbook.open(workbook_path)
    assert apply_jobs(fresh, jobs) == []
    assert jobs[0].status == "error"


def test_due_status(workbook_path):
    wb = ToolWorkbook.open(workbook_path)
    rows = due_status(wb.records, today=dt.date(2026, 7, 1), warn_days=30)
    states = {r["Serial"]: r["상태"] for r in rows}
    assert states["0716805829"] == "만료"  # due 2026-06-24
    assert states["0116503968"] == "폐기"
    assert states["31550075"] == "정상"
    assert rows[0]["상태"] == "만료"


def test_stale_smax_screenshot_is_flagged(workbook_path):
    wb = ToolWorkbook.open(workbook_path)
    old_due = dt.date(2026, 12, 29)  # what SMAX still shows
    recognizer = StubRecognizer({
        900: _tool("박민수", "31550075"),
        1500: Recognition(image_type="smax_screenshot", serial_number="31550075", next_due_date=old_due.isoformat()),
    })
    smax = jpeg_bytes(Image.new("RGB", (1500, 450), (255, 255, 255)))
    jobs = analyze_photos(wb, [("a.jpg", _photo(900)), ("s.png", smax)], recognizer)
    assert any("SMAX" in m and "다릅니다" in m for m in jobs[1].messages)


def test_sideways_photo_is_turned_upright(workbook_path):
    wb = ToolWorkbook.open(workbook_path)
    upright = make_tool_photo(size=(900, 1200))
    sideways = upright.rotate(90, expand=True)  # 1200x900, text runs bottom-to-top
    reading = _tool("박민수", "31550075")
    reading.rotation = 90
    jobs = analyze_photos(wb, [("s.jpg", jpeg_bytes(sideways))], StubRecognizer({1200: reading}))
    job = jobs[0]
    assert job.image.size == (900, 1200)
    w, h = job.label_crop.size
    assert w > h  # the label reads left-to-right again
