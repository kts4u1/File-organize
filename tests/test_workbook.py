import datetime as dt
import zipfile

import openpyxl
from lxml import etree

from conftest import build_workbook, jpeg_bytes
from PIL import Image
from toolcal.workbook import ToolWorkbook
from toolcal.xlsx_package import NS


JPEG = jpeg_bytes(Image.new("RGB", (40, 30), (250, 250, 250)))


def _record(wb, sheet, serial):
    return next(r for r in wb.records if r.sheet == sheet and r.serial == serial)


def test_reads_records(workbook_path):
    wb = ToolWorkbook.open(workbook_path)
    assert set(wb.sheets) == {"대전", "광주"}
    assert len(wb.records) == 7
    rec = _record(wb, "광주", "31550075")
    assert rec.row == 3
    assert rec.cal_last == dt.date(2025, 12, 29)
    assert rec.cal_next == dt.date(2026, 12, 29)
    assert rec.interval_months == 12
    assert rec.names == ["박민수"]
    assert rec.picture_counts == {"label_photo": 1, "smax_photo": 1}
    # numeric serials come back as text, leading zeros are kept for text cells
    assert _record(wb, "대전", "218800584").interval_months == 6
    disposed = _record(wb, "광주", "0116503968")
    assert disposed.disposed and disposed.names == ["박민수"]


def test_write_dates_and_picture(workbook_path, tmp_path):
    wb = ToolWorkbook.open(workbook_path)
    rec = _record(wb, "광주", "31550075")
    assert wb.write_date(rec, "cal_last", dt.date(2026, 10, 10)) == "written"
    assert wb.write_date(rec, "cal_next", dt.date(2027, 10, 10)) == "written"
    assert wb.write_date(rec, "cal_next", dt.date(2027, 10, 10)) == "unchanged"
    removed = wb.replace_picture(rec, "label_photo", JPEG, "jpeg", (400, 300))
    assert removed == 1
    out = tmp_path / "out.xlsx"
    wb.save(out)

    xl = openpyxl.load_workbook(out)
    ws = xl["광주"]
    assert ws["B3"].value == dt.datetime(2026, 10, 10)
    assert ws["A3"].value == dt.datetime(2027, 10, 10)
    assert ws["A3"].is_date
    assert ws["A2"].value == dt.datetime(2026, 12, 29)  # neighbours untouched

    again = ToolWorkbook.open(out)
    rec2 = _record(again, "광주", "31550075")
    assert rec2.picture_counts == {"label_photo": 1, "smax_photo": 1}
    assert rec2.cal_next == dt.date(2027, 10, 10)


def test_only_touched_parts_change(workbook_path, tmp_path):
    wb = ToolWorkbook.open(workbook_path)
    rec = _record(wb, "광주", "31550075")
    wb.write_date(rec, "cal_last", dt.date(2026, 10, 10))
    wb.replace_picture(rec, "label_photo", JPEG, "jpeg", (400, 300))
    out = tmp_path / "out.xlsx"
    wb.save(out)
    before, after = zipfile.ZipFile(workbook_path), zipfile.ZipFile(out)
    changed = {n for n in set(before.namelist()) & set(after.namelist()) if before.read(n) != after.read(n)}
    gwangju_sheet = ToolWorkbook.open(workbook_path).sheets["광주"].part
    drawing_parts = {n for n in changed if n.startswith("xl/drawings/")}
    assert changed - drawing_parts == {"[Content_Types].xml", gwangju_sheet}
    assert len(drawing_parts) == 2  # the 광주 drawing and its .rels
    added = set(after.namelist()) - set(before.namelist())
    removed = set(before.namelist()) - set(after.namelist())
    assert len(added) == 1 and added.pop().endswith(".jpeg")
    assert len(removed) == 1 and removed.pop().startswith("xl/media/")
    assert after.namelist()[0] == "[Content_Types].xml"


def test_picture_anchor_fits_cell(workbook_path, tmp_path):
    wb = ToolWorkbook.open(workbook_path)
    rec = _record(wb, "대전", "46900004")
    wb.replace_picture(rec, "label_photo", JPEG, "jpeg", (1000, 700))
    info = wb.sheets["대전"]
    drawing = wb._drawing_part(info, create=False)
    anchors = wb._anchors_at(info, rec.row, "J")
    assert len(anchors) == 1
    a = anchors[0]
    frm, to = a.find(f"{{{NS['xdr']}}}from"), a.find(f"{{{NS['xdr']}}}to")
    assert frm.findtext(f"{{{NS['xdr']}}}col") == to.findtext(f"{{{NS['xdr']}}}col") == "9"
    assert frm.findtext(f"{{{NS['xdr']}}}row") == to.findtext(f"{{{NS['xdr']}}}row") == str(rec.row - 1)
    width_px = int(to.findtext(f"{{{NS['xdr']}}}colOff")) / 9525
    height_px = int(to.findtext(f"{{{NS['xdr']}}}rowOff")) / 9525
    assert width_px <= wb._col_width_px(info, 10)
    assert height_px <= 100 * 96 / 72
    ext = a.find(f".//{{{NS['a']}}}ext")
    assert abs(int(ext.get("cx")) / int(ext.get("cy")) - 1000 / 700) < 0.02
    # shape ids stay unique
    ids = [el.get("id") for el in wb.pkg.xml(drawing).iter(f"{{{NS['xdr']}}}cNvPr")]
    assert len(ids) == len(set(ids))


def test_formula_cells_are_left_alone(tmp_path):
    path = build_workbook(tmp_path / "f.xlsx", with_pictures=False)
    xl = openpyxl.load_workbook(path)
    ws = xl["대전"]
    ws["A2"] = "=EDATE(B2,12)"
    xl.save(path)
    wb = ToolWorkbook.open(path)
    rec = next(r for r in wb.records if r.sheet == "대전" and r.row == 2)
    assert wb.write_date(rec, "cal_last", dt.date(2026, 10, 1)) == "written"
    assert wb.write_date(rec, "cal_next", dt.date(2027, 10, 1)) == "formula"
    wb.save(path)
    xl = openpyxl.load_workbook(path)
    assert xl["대전"]["A2"].value == "=EDATE(B2,12)"
    calc = etree.fromstring(zipfile.ZipFile(path).read("xl/workbook.xml")).find(f"{{{NS['main']}}}calcPr")
    assert calc.get("fullCalcOnLoad") == "1"


def test_text_cell_gets_date_style(workbook_path, tmp_path):
    wb = ToolWorkbook.open(workbook_path)
    rec = _record(wb, "광주", "0116503968")  # Cal Next holds the text '폐기'
    wb.write_date(rec, "cal_next", dt.date(2026, 12, 27))
    out = tmp_path / "out.xlsx"
    wb.save(out)
    cell = openpyxl.load_workbook(out)["광주"]["A5"]
    assert cell.value == dt.datetime(2026, 12, 27) and cell.is_date


def test_sheet_without_drawing_gets_one(tmp_path):
    path = build_workbook(tmp_path / "d.xlsx", with_pictures=False, third_sheet_without_drawing=True)
    wb = ToolWorkbook.open(path)
    rec = next(r for r in wb.records if r.sheet == "부산")
    assert wb.replace_picture(rec, "label_photo", JPEG, "jpeg", (300, 200)) == 0
    out = tmp_path / "out.xlsx"
    wb.save(out)
    xl = openpyxl.load_workbook(out)
    assert len(xl["부산"]._images) == 1
    again = ToolWorkbook.open(out)
    assert next(r for r in again.records if r.sheet == "부산").picture_counts["label_photo"] == 1
